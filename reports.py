# comprehensive_deriv_analysis.py - FIXED VERSION
import MetaTrader5 as mt5
import pandas as pd
import numpy as np
from datetime import datetime, timedelta
import os
from pathlib import Path
import matplotlib.pyplot as plt
import seaborn as sns
import re
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import PieChart, BarChart, Reference
import matplotlib.font_manager as fm
from typing import List, Dict, Any, Tuple
import shutil

# === CONFIGURATION ===
MT5_PATH = r"C:\Program Files\MetaTrader 5 Terminal\terminal64.exe"
OUTPUT_FOLDER = str(Path.home() / "OneDrive - University of Ghana" / "Deriv_Comprehensive_Analysis")
os.makedirs(OUTPUT_FOLDER, exist_ok=True)



def clear_output_folder():
    """Delete all files and subfolders inside OUTPUT_FOLDER, but keep the folder itself."""
    if os.path.exists(OUTPUT_FOLDER):
        for item in os.listdir(OUTPUT_FOLDER):
            item_path = os.path.join(OUTPUT_FOLDER, item)
            try:
                if os.path.isfile(item_path) or os.path.islink(item_path):
                    os.unlink(item_path)        # delete file or link
                elif os.path.isdir(item_path):
                    shutil.rmtree(item_path)    # delete folder and everything inside
                print(f"Deleted: {item_path}")
            except Exception as e:
                print(f"Warning: Failed to delete {item_path}: {e}")
        print(f"Output folder cleared: {OUTPUT_FOLDER}")
    else:
        os.makedirs(OUTPUT_FOLDER, exist_ok=True)
        print(f"Output folder created: {OUTPUT_FOLDER}")

plt.style.use('seaborn-v0_8-whitegrid')
sns.set_palette("husl")
plt.rcParams['figure.figsize'] = (12, 8)
plt.rcParams['font.size'] = 10

# Connect to MT5
if not mt5.initialize(path=MT5_PATH):
    print("❌ Failed to connect to MT5")
    quit()
print("✅ Connected to MT5")

# Timeframe mapping
TIMEFRAMES = {
    1: "M1", 5: "M5", 15: "M15", 30: "M30",
    16385: "H1", 16388: "H4", 16408: "D1"
}


# Trading configuration from visuals.py
def get_trading_config():
    FOREX_PAIRS = ["EURUSD", "GBPUSD", "USDJPY", "CHFJPY",  "USDCHF","EURJPY"]
    XAUUSD = "XAUUSD"
    BTCUSD="BTCUSD"
    MAJORS = [ "EURCHF", "CADCHF","GBPCAD","USDCAD"]
    V100 = "Volatility 100 Index"
    BnC = ["Boom 900 Index","Boom 300 Index"]
    Jump = ["Jump 25 Index"]
    vol = ["Volatility 50 Index","Volatility 25 Index","Volatility 10 Index"]
    SYNTHETICS = ["Step Index"]
    US =["US Tech 100","Wall Street 30"]
    v=["Volatility 75 Index"]
    J=["Jump 75 Index"]
    C=["Crash 900 Index","Crash 300 Index"]
    return {
        "H1_PENDING": MAJORS+Jump+[BTCUSD]+J ,
        "H1_INSTANT": [XAUUSD]+US+SYNTHETICS+C+BnC,
        "M15_PENDING": Jump+FOREX_PAIRS+vol+[XAUUSD,V100]+SYNTHETICS+v,
        "M15_INSTANT": J+US,
        "M5_PENDING": BnC+MAJORS+FOREX_PAIRS+vol+[XAUUSD,V100,BTCUSD]+SYNTHETICS+Jump+C+J+US,
        "M5_INSTANT": []
    }

def get_all_deals():
    from_date = datetime(2025, 10, 1)
    to_date = datetime.now()
    deals = mt5.history_deals_get(from_date, to_date)
    print(f"Looking for deals from {from_date.date()} to {to_date.date()}")
    if deals is None:
        print("history_deals_get returned None")
    else:
        print(f"Found {len(deals)} total deals in history")
        if len(deals) > 0:
            print(f"Most recent deal time: {pd.to_datetime(deals[-1].time, unit='s')}")
    if deals is None or len(deals) == 0:
        print("⚠️ No deals found")
        return pd.DataFrame()
    df = pd.DataFrame(list(deals), columns=deals[0]._asdict().keys())
    if 'time' in df.columns:
        df['time'] = pd.to_datetime(df['time'], unit='s')
    return df

def get_all_orders():
    from_date = datetime(2025, 10, 1)
    to_date = datetime.now()
    orders = mt5.history_orders_get(from_date, to_date)
    if orders is None or len(orders) == 0:
        print("⚠️ No orders found")
        return pd.DataFrame()
    df = pd.DataFrame(list(orders), columns=orders[0]._asdict().keys())
    if 'time_setup' in df.columns:
        df['time_setup'] = pd.to_datetime(df['time_setup'], unit='s')
    if 'time_done' in df.columns:
        df['time_done'] = pd.to_datetime(df['time_done'], unit='s')
    return df

def extract_tp1_from_comment(comment: str):
    """Extract TP1 from comment field"""
    if not comment or pd.isna(comment):
        return None
    try:
        numbers = re.findall(r"[-+]?\d*\.\d+|\d+", str(comment))
        if numbers:
            return float(numbers[0])
    except (ValueError, TypeError):
        pass
    return None

def safe_float_conversion(value, default=0.0):
    """Safely convert value to float"""
    try:
        return float(value) if value is not None and value != "" else default
    except (ValueError, TypeError):
        return default

def analyze_pending_orders_enhanced():
    """Enhanced pending order analysis with better edge case handling"""
    print("🔍 Enhanced Pending Order Analysis...")
    orders_df = get_all_orders()
    if orders_df.empty:
        print("No historical orders found.")
        return []

    # Filter pending orders
    pending = orders_df[orders_df['type'].isin([2, 3, 4, 5])].copy()
    print(f"Found {len(pending)} pending orders to analyze")

    records = []
    config = get_trading_config()

    for _, order in pending.iterrows():
        symbol = order['symbol']
        magic = order['magic']
        tf_str = TIMEFRAMES.get(magic, f"TF{magic}")
        direction = "BUY" if order['type'] in [2, 4] else "SELL"
        
        # Safe conversion
        entry = safe_float_conversion(order['price_open'])
        sl = safe_float_conversion(order['sl'])
        tp2 = safe_float_conversion(order['tp'])
        comment = str(order.get('comment', ''))
        tp1 = extract_tp1_from_comment(comment)

        if tp1 is None or sl == 0.0 or entry == 0.0:
            continue

        start = order['time_setup']
        end = order['time_done'] if pd.notna(order['time_done']) else datetime.now()
        
        # More realistic end time calculation
        if pd.notna(order['time_done']):
            # Order was completed/cancelled
            position_id = order.get('position_id', 0)
            if position_id != 0:
                # This became a position, use position logic
                deals = mt5.history_deals_get(position=position_id)
                if deals and len(deals) > 1:
                    # Use the actual exit time
                    exit_time = max(pd.to_datetime(deal.time, unit='s') for deal in deals)
                    end = exit_time
            # If cancelled without becoming position, keep the cancellation time
        else:
            # Still pending, analyze up to current time
            end = datetime.now()

        # Get price data with validation
        try:
            rates = mt5.copy_rates_range(symbol, magic, start, end)
            if rates is None or len(rates) == 0:
                print(f"⚠️ No price data for {symbol} in range {start} to {end}")
                continue

            df_rates = pd.DataFrame(rates)
            df_rates['time'] = pd.to_datetime(df_rates['time'], unit='s')
            
            # Check if we have sufficient data
            if len(df_rates) < 2:
                print(f"⚠️ Insufficient price data for {symbol}")
                continue

            # Track events with improved logic
            tp1_hit = sl_hit = entry_hit = False
            tp1_time = sl_time = entry_time = None
            first_touch_time = None

            for _, candle in df_rates.iterrows():
                h = safe_float_conversion(candle['high'])
                l = safe_float_conversion(candle['low'])
                candle_time = candle['time']

                # Track which level was hit first
                if not first_touch_time:
                    if ((direction == "BUY" and (h >= tp1 or l <= sl or l <= entry)) or 
                        (direction == "SELL" and (l <= tp1 or h >= sl or h >= entry))):
                        first_touch_time = candle_time

                # TP1 check
                if not tp1_hit:
                    if (direction == "BUY" and h >= tp1) or (direction == "SELL" and l <= tp1):
                        tp1_hit = True
                        tp1_time = candle_time

                # SL check
                if not sl_hit and sl > 0:
                    if (direction == "BUY" and l <= sl) or (direction == "SELL" and h >= sl):
                        sl_hit = True
                        sl_time = candle_time

                # Entry check
                if not entry_hit:
                    if (direction == "BUY" and l <= entry) or (direction == "SELL" and h >= entry):
                        entry_hit = True
                        entry_time = candle_time

            # Enhanced outcome determination
            missed_opportunity = tp1_hit and not entry_hit
            won_trade = False
            lost_trade = False
            
            if entry_hit:
                if tp1_hit and (not sl_hit or (sl_time and tp1_time and tp1_time < sl_time)):
                    won_trade = True
                elif sl_hit and (not tp1_hit or (tp1_time and sl_time and sl_time < tp1_time)):
                    lost_trade = True

            # Determine order type from config
            order_type = "UNKNOWN"
            if tf_str == "H1":
                if symbol in config["H1_PENDING"]: order_type = "PENDING"
                elif symbol in config["H1_INSTANT"]: order_type = "INSTANT"
            elif tf_str == "M15":
                if symbol in config["M15_PENDING"]: order_type = "PENDING"
                elif symbol in config["M15_INSTANT"]: order_type = "INSTANT"
            elif tf_str == "M5":
                if symbol in config["M5_PENDING"]: order_type = "PENDING"
                elif symbol in config["M5_INSTANT"]: order_type = "INSTANT"

            records.append({
                "Symbol": symbol,
                "Timeframe": tf_str,
                "Direction": direction,
                "Order Type": order_type,
                "Setup Time": start,
                "Entry Price": entry,
                "TP1": tp1,
                "SL": sl,
                "TP2": tp2,
                "TP1 Hit": tp1_hit,
                "SL Hit": sl_hit,
                "Entry Filled": entry_hit,
                "Missed Opportunity": missed_opportunity,
                "Won Trade": won_trade,
                "Lost Trade": lost_trade,
                "TP1 Time": tp1_time,
                "SL Time": sl_time,
                "Entry Time": entry_time,
                "First Touch Time": first_touch_time,
                "Analysis Period": f"{(end - start).days} days",
                "Trade Type": "PENDING_ORDER"
            })
            
        except Exception as e:
            print(f"⚠️ Error analyzing {symbol}: {e}")
            continue

    print(f"Enhanced pending order analysis complete: {len(records)} orders processed")
    return records

def analyze_completed_trades():
    """Analyze completed trades from deals history - ROBUST VERSION"""
    print("Analyzing completed trades...")
    deals = get_all_deals()
    records = []
    
    if deals.empty:
        print("No deals found")
        return records

    # Group by position_id
    for pos_id, group in deals.groupby('position_id'):
        if pos_id == 0:
            continue
        
        # Find real entry deal (type 0 or 1, entry=0)
        entry_deals = group[group['entry'] == 0]
        if entry_deals.empty:
            continue
            
        entry = entry_deals.iloc[0]
        
        # === SAFELY extract fields (this is the fix) ===
        symbol = entry.get('symbol', 'Unknown')
        magic = int(entry.get('magic', 0)) if entry.get('magic') is not None else 0
        timeframe_str = TIMEFRAMES.get(magic, f"Unknown({magic})")
        
        direction = "BUY" if entry['type'] == 0 else "SELL"
        volume = safe_float_conversion(entry.get('volume', 0))
        entry_price = safe_float_conversion(entry.get('price', 0))
        entry_time = pd.to_datetime(entry['time'], unit='s')
        
        sl = safe_float_conversion(entry.get('sl', 0))
        tp = safe_float_conversion(entry.get('tp', 0))
        
        comment = str(entry.get('comment', '')).strip()
        tp1 = extract_tp1_from_comment(comment)
        
        net_profit = group['profit'].sum() + group['swap'].sum() + group['commission'].sum()
        net_profit = round(float(net_profit), 2)
        
        # Outcome & Reason
        if net_profit > 0:
            outcome = "WINNER"
            reason = "TP Hit"
        elif net_profit < 0:
            outcome = "LOSER"
            reason = "SL Hit"
        else:
            outcome = "BREAKEVEN"
            reason = "Manual/BE"
        
        # Order type detection
        config = get_trading_config()
        order_type = "UNKNOWN"
        tf_key = None
        if magic == 16385 or timeframe_str == "H1":
            tf_key = "H1"
        elif magic == mt5.TIMEFRAME_M15 or timeframe_str == "M15":
            tf_key = "M15"
        elif magic == mt5.TIMEFRAME_M5 or timeframe_str == "M5":
            tf_key = "M5"
            
        if tf_key:
            if symbol in config.get(f"{tf_key}_PENDING", []) or symbol in config.get(f"{tf_key}_INSTANT", []):
                order_type = "PENDING" if symbol in config.get(f"{tf_key}_PENDING", []) else "INSTANT"

        records.append({
            "Symbol": symbol,
            "Timeframe": timeframe_str,
            "Direction": direction,
            "Order Type": order_type,
            "Entry Time": entry_time,
            "Profit": net_profit,
            "Outcome": outcome,
            "Reason": reason,
            "Trade Type": "COMPLETED",
            "Entry Price": entry_price,
            "SL": sl if sl != 0 else None,
            "TP": tp if tp != 0 else None,
            "TP1": tp1,
            "Volume": volume,
            "Position ID": pos_id
        })

    print(f"Completed trades analysis: {len(records)} trades processed")
    return records

def analyze_strategy_performance():
    """Comprehensive analysis combining pending orders and completed trades - FINAL FIX"""
    print("Running comprehensive strategy analysis...")
    
    pending_data = analyze_pending_orders_enhanced()
    completed_data = analyze_completed_trades()
    
    # === THIS IS THE REAL FIX ===
    # Convert separately first so columns are preserved
    df_pending   = pd.DataFrame(pending_data)
    df_completed = pd.DataFrame(completed_data)
    
    # Add missing columns with correct type so they survive the concat
    for col in ['Outcome', 'Reason', 'TP', 'Volume', 'Position ID', 'Profit']:
        if col not in df_pending.columns:
            df_pending[col] = pd.NA
        if col not in df_completed.columns:
            df_completed[col] = pd.NA
    
    # Now safe concat
    df = pd.concat([df_pending, df_completed], ignore_index=True)
    # ==================================
    
    # Type fixing (same as before)
    numeric_columns = ['Entry Price', 'TP1', 'SL', 'TP2', 'Profit', 'Volume', 'TP']
    for col in numeric_columns:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors='coerce')
    
    bool_columns = ['TP1 Hit', 'SL Hit', 'Entry Filled', 'Missed Opportunity', 'Won Trade']
    for col in bool_columns:
        if col in df.columns:
            df[col] = df[col].astype('boolean')
    
    print(f"Final combined dataset: {len(df)} rows ({len(df_pending)} pending + {len(df_completed)} completed)")
    return df

def generate_comprehensive_report(df: pd.DataFrame):
    """Generate comprehensive analysis report"""
    if df.empty:
        print("No data for report generation")
        return
    
    # Separate data by type
    pending_orders = df[df['Trade Type'] == 'PENDING_ORDER']
    completed_trades = df[df['Trade Type'] == 'COMPLETED']
    
    print(f"\n{'='*80}")
    print("🎯 COMPREHENSIVE TRADING BOT ANALYSIS REPORT")
    print(f"{'='*80}")
    
    # 1. PENDING ORDER ANALYSIS
    if not pending_orders.empty:
        print(f"\n📊 PENDING ORDER ANALYSIS ({len(pending_orders)} orders)")
        print(f"{'-'*60}")
        
        # Missed opportunities analysis
        missed_opps = pending_orders[pending_orders['Missed Opportunity'] == True]
        missed_by_symbol_tf = missed_opps.groupby(['Symbol', 'Timeframe']).size().reset_index(name='Missed Count')
        missed_by_symbol_tf = missed_by_symbol_tf.sort_values('Missed Count', ascending=False)
        
        print(f"📈 Missed Opportunities (TP1 hit but no entry): {len(missed_opps)}")
        if not missed_opps.empty:
            print("\nTop missed opportunities:")
            for _, row in missed_by_symbol_tf.head(10).iterrows():
                print(f"   {row['Symbol']} - {row['Timeframe']}: {row['Missed Count']} missed")
        
        # Execution rate by symbol+timeframe - FIXED VERSION
        exec_stats = pending_orders.groupby(['Symbol', 'Timeframe']).agg({
            'Entry Filled': ['count', 'sum']
        }).reset_index()
        
        # Flatten column names
        exec_stats.columns = ['Symbol', 'Timeframe', 'Total_Orders', 'Executed_Orders']
        
        # Convert to numeric and calculate execution rate
        exec_stats['Total_Orders'] = pd.to_numeric(exec_stats['Total_Orders'], errors='coerce')
        exec_stats['Executed_Orders'] = pd.to_numeric(exec_stats['Executed_Orders'], errors='coerce')
        exec_stats['Execution_Rate_%'] = (exec_stats['Executed_Orders'] / exec_stats['Total_Orders'] * 100).round(1)
        exec_stats = exec_stats.sort_values('Execution_Rate_%')
        
        print(f"\n📉 Lowest Execution Rates:")
        for _, row in exec_stats.head(5).iterrows():
            print(f"   {row['Symbol']} - {row['Timeframe']}: {row['Execution_Rate_%']}% ({row['Executed_Orders']:.0f}/{row['Total_Orders']:.0f})")
    
    # 2. COMPLETED TRADES ANALYSIS
    if not completed_trades.empty:
        print(f"\n💰 COMPLETED TRADES ANALYSIS ({len(completed_trades)} trades)")
        print(f"{'-'*60}")
        
        # Basic stats with safe numeric conversion
        total_profit = pd.to_numeric(completed_trades['Profit'], errors='coerce').sum()
        win_rate = (completed_trades['Outcome'] == 'WINNER').mean() * 100
        avg_profit = pd.to_numeric(completed_trades['Profit'], errors='coerce').mean()
        
        print(f"Total Profit: ${total_profit:.2f}")
        print(f"Win Rate: {win_rate:.1f}%")
        print(f"Average Profit per Trade: ${avg_profit:.2f}")
        
        # Performance by Symbol+Timeframe+Direction
        stfd_performance = completed_trades.groupby(['Symbol', 'Timeframe', 'Direction']).agg({
            'Profit': ['count', 'sum', 'mean'],
        }).round(2)
        
        if not stfd_performance.empty:
            # Flatten column names
            stfd_performance.columns = ['Trade_Count', 'Total_Profit', 'Avg_Profit']
            
            # Calculate win rate separately
            win_rates = completed_trades.groupby(['Symbol', 'Timeframe', 'Direction'])['Outcome'].apply(
                lambda x: (x == 'WINNER').mean() * 100
            ).round(1)
            
            # Combine with main stats
            stfd_performance['Win_Rate_%'] = win_rates
            stfd_performance = stfd_performance.sort_values('Total_Profit', ascending=False)
            
            print(f"\n🏆 TOP PERFORMING COMBINATIONS:")
            top_combinations = stfd_performance.head(5)
            for idx, stats in top_combinations.iterrows():
                symbol, tf, direction = idx
                print(f"   ✅ {symbol} + {tf} + {direction}: {stats['Win_Rate_%']}% WR, ${stats['Total_Profit']:.2f} profit")
            
            # Filter for combinations with sufficient trades
            sufficient_trades = stfd_performance[stfd_performance['Trade_Count'] >= 3]
            if not sufficient_trades.empty:
                print(f"\n📉 WORST PERFORMING COMBINATIONS (min 3 trades):")
                worst_combinations = sufficient_trades.tail(5)
                for idx, stats in worst_combinations.iterrows():
                    symbol, tf, direction = idx
                    print(f"   ❌ {symbol} + {tf} + {direction}: {stats['Win_Rate_%']}% WR, ${stats['Total_Profit']:.2f} profit")
    
    # 3. STRATEGY RECOMMENDATIONS
    print(f"\n🎯 STRATEGY OPTIMIZATION RECOMMENDATIONS")
    print(f"{'-'*60}")
    
    # Pending vs Instant recommendations
    if not pending_orders.empty:
        high_missed = pending_orders[pending_orders['Missed Opportunity'] == True]
        high_missed_grouped = high_missed.groupby(['Symbol', 'Timeframe']).size()
        high_missed_grouped = high_missed_grouped[high_missed_grouped >= 3]
        
        if not high_missed_grouped.empty:
            print("🔔 CONSIDER SWITCHING TO INSTANT ORDERS:")
            for (symbol, timeframe), count in high_missed_grouped.items():
                print(f"   📍 {symbol} on {timeframe}: {count} missed TP1 hits")
    
    # Risk management recommendations
    if not completed_trades.empty:
        losing_trades = completed_trades[completed_trades['Profit'] < 0]
        if not losing_trades.empty:
            avg_loss = pd.to_numeric(losing_trades['Profit'], errors='coerce').mean()
            max_loss = pd.to_numeric(losing_trades['Profit'], errors='coerce').min()
            print(f"📉 Risk Management:")
            print(f"   Average Loss: ${avg_loss:.2f}")
            print(f"   Maximum Loss: ${max_loss:.2f}")
            
            if max_loss < -100:
                print(f"   ⚠️  Consider tighter SL for large losses")

def create_enhanced_visualizations(df: pd.DataFrame):
    """Create comprehensive visualizations — NOW WITH DIRECTION IN MISSED OPPORTUNITIES"""
    if df.empty:
        print("No data for visualizations")
        return
    
    viz_folder = os.path.join(OUTPUT_FOLDER, "visualizations")
    os.makedirs(viz_folder, exist_ok=True)
    
    # Separate data
    pending_orders = df[df['Trade Type'] == 'PENDING_ORDER']
    completed_trades = df[df['Trade Type'] == 'COMPLETED']
    
    # 1. MISSED OPPORTUNITIES CHART — NOW WITH DIRECTION (BUY/SELL)
    if not pending_orders.empty:
        missed_opps = pending_orders[pending_orders['Missed Opportunity'] == True]
        if not missed_opps.empty:
            # Group by Symbol + Timeframe + Direction
            missed_by_group = missed_opps.groupby(['Symbol', 'Timeframe', 'Direction']).size().reset_index(name='Count')
            missed_by_group = missed_by_group.sort_values('Count', ascending=False)
            
            # Create detailed label
            missed_by_group['Label'] = (
                missed_by_group['Symbol'] + ' | ' +
                missed_by_group['Timeframe'] + ' | ' +
                missed_by_group['Direction']
            )
            
            plt.figure(figsize=(16, 9))
            
            # Color mapping: BUY = Green, SELL = Red
            colors = ['#27ae60' if direction == 'BUY' else '#c0392b' for direction in missed_by_group['Direction']]
            
            bars = plt.bar(missed_by_group['Label'], missed_by_group['Count'], 
                          color=colors, edgecolor='black', linewidth=1.2, alpha=0.9)
            
            plt.title('Missed Opportunities (TP1 Hit Before Entry)\nGrouped by Symbol • Timeframe • Direction', 
                     fontsize=18, fontweight='bold', pad=20)
            plt.ylabel('Number of Missed Trades', fontsize=14, fontweight='bold')
            plt.xlabel('Symbol | Timeframe | Direction', fontsize=14, fontweight='bold')
            plt.xticks(rotation=50, ha='right', fontsize=10)
            plt.grid(axis='y', alpha=0.3, linestyle='--')
            
            # Add value labels on top of bars
            for bar, count, direction in zip(bars, missed_by_group['Count'], missed_by_group['Direction']):
                plt.text(bar.get_x() + bar.get_width()/2, bar.get_height() + 0.15,
                        f'{int(count)}', ha='center', va='bottom', fontweight='bold', fontsize=11)
            
            # Add legend
            from matplotlib.patches import Patch
            legend_elements = [
                Patch(facecolor='#27ae60', edgecolor='black', label='BUY Signals Missed'),
                Patch(facecolor='#c0392b', edgecolor='black', label='SELL Signals Missed')
            ]
            plt.legend(handles=legend_elements, loc='upper right', fontsize=12)
            
            plt.tight_layout()
            plt.savefig(os.path.join(viz_folder, 'missed_opportunities_with_direction.png'), 
                       dpi=300, bbox_inches='tight', facecolor='white')
            plt.close()
            
            print(f"Created enhanced missed opportunities chart with DIRECTION: {len(missed_opps)} missed trades")

    # 2. Performance by Symbol+Timeframe+Direction (unchanged — already includes direction)
    if not completed_trades.empty:
        stfd_performance = completed_trades.groupby(['Symbol', 'Timeframe', 'Direction']).agg({
            'Profit': ['count', 'sum', 'mean'],
        }).round(2)
        
        if not stfd_performance.empty:
            stfd_performance.columns = ['Trade_Count', 'Total_Profit', 'Avg_Profit']
            
            win_rates = completed_trades.groupby(['Symbol', 'Timeframe', 'Direction'])['Outcome'].apply(
                lambda x: (x == 'WINNER').mean() * 100
            ).round(1)
            stfd_performance['Win_Rate_%'] = win_rates
            stfd_performance = stfd_performance.sort_values('Total_Profit', ascending=False)
            
            top_combinations = stfd_performance.head(15)
            
            if not top_combinations.empty:
                fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(16, 12))
                fig.suptitle('TOP PERFORMING SYMBOL + TIMEFRAME + DIRECTION COMBINATIONS', 
                            fontsize=18, fontweight='bold', y=0.96)
                
                labels = [f"{idx[0]}\n{idx[1]}\n{idx[2]}" for idx in top_combinations.index]
                profit_colors = ['#2E8B57' if x > 0 else '#DC143C' for x in top_combinations['Total_Profit']]
                
                # Profit
                bars1 = ax1.bar(range(len(top_combinations)), top_combinations['Total_Profit'], 
                              color=profit_colors, alpha=0.85, edgecolor='black')
                ax1.set_title('Total Profit ($)', fontsize=14, fontweight='bold')
                ax1.set_ylabel('Profit ($)', fontweight='bold')
                ax1.set_xticks(range(len(top_combinations)))
                ax1.set_xticklabels(labels, rotation=45, ha='right', fontsize=9)
                ax1.grid(axis='y', alpha=0.3)
                
                for bar, val in zip(bars1, top_combinations['Total_Profit']):
                    ax1.text(bar.get_x() + bar.get_width()/2, bar.get_height() + (8 if val > 0 else -20),
                            f'${val:.0f}', ha='center', va='bottom' if val > 0 else 'top',
                            fontweight='bold', color='black', fontsize=9)
                
                # Win Rate
                bars2 = ax2.bar(range(len(top_combinations)), top_combinations['Win_Rate_%'], 
                              color='#4682B4', alpha=0.8, edgecolor='black')
                ax2.set_title('Win Rate (%)', fontsize=14, fontweight='bold')
                ax2.set_ylabel('Win Rate (%)', fontweight='bold')
                ax2.set_xlabel('Symbol • Timeframe • Direction', fontweight='bold')
                ax2.set_xticks(range(len(top_combinations)))
                ax2.set_xticklabels(labels, rotation=45, ha='right', fontsize=9)
                ax2.set_ylim(0, 100)
                ax2.grid(axis='y', alpha=0.3)
                
                for bar, wr in zip(bars2, top_combinations['Win_Rate_%']):
                    ax2.text(bar.get_x() + bar.get_width()/2, bar.get_height() + 2,
                            f'{wr:.0f}%', ha='center', va='bottom', fontweight='bold', fontsize=10)
                
                plt.tight_layout(rect=[0, 0, 1, 0.94])
                plt.savefig(os.path.join(viz_folder, 'performance_analysis.png'), 
                           dpi=300, bbox_inches='tight', facecolor='white')
                plt.close()
                print("Created performance analysis chart")

def save_comprehensive_excel_report(df: pd.DataFrame):
    """Save comprehensive analysis to Excel with formatting — NOW INCLUDES DIRECTION"""
    if df.empty:
        print("No data to save")
        return
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    filename = os.path.join(OUTPUT_FOLDER, f"Deriv_Comprehensive_Analysis_{timestamp}.xlsx")
    
    try:
        with pd.ExcelWriter(filename, engine='openpyxl') as writer:
            # Main data sheet
            df.to_excel(writer, sheet_name="All_Data", index=False)
            
            # === 1. PENDING ORDERS SUMMARY (NOW WITH DIRECTION) ===
                        # === 1. PENDING ORDERS SUMMARY – NEW METRICS YOU ASKED FOR ===
                        # === 1. PENDING ORDERS SUMMARY – FIXED & CLEAN VERSION ===
            pending_orders = df[df['Trade Type'] == 'PENDING_ORDER']
            if not pending_orders.empty:
                def pending_stats(group):
                    total = len(group)
                    filled = group['Entry Filled'].sum()
                    missed = group['Missed Opportunity'].sum()
                    won_after_fill = group['Won Trade'].sum()

                    won_executed_pct = (won_after_fill / filled * 100) if filled > 0 else 0
                    missed_rate_pct = (missed / total * 100)

                    return pd.Series({
                        'Total_Orders': total,
                        'Filled_Orders': int(filled),
                        'Missed_Opportunities': int(missed),
                        'Won_After_Fill': int(won_after_fill),
                        'Won_Executed_%': round(won_executed_pct, 1),
                        'Lost_After_Fill': int(filled - won_after_fill),
                        'Missed_Opportunity_Rate_%': round(missed_rate_pct, 1)
                    })

                # Fix: explicitly include grouping columns + suppress warning
                pending_summary = (
                    pending_orders
                    .groupby(['Symbol', 'Timeframe', 'Direction', 'Order Type'], as_index=False)
                    .apply(pending_stats, include_groups=False)  # this removes the warning
                )

                pending_summary = pending_summary.sort_values(
                    by=['Symbol', 'Timeframe', 'Direction', 'Order Type', 'Missed_Opportunities'],
                    ascending=[True, True, True, True, False]
                ).reset_index(drop=True)
                pending_summary.to_excel(writer, sheet_name="Pending_Orders_Summary", index=False)

            # === 2. COMPLETED TRADES SUMMARY (NOW WITH DIRECTION) ===
                        # === 2. COMPLETED TRADES SUMMARY – NEW VERSION WITH WINNERS/LOSERS ===
                        # === 2. COMPLETED TRADES SUMMARY – FIXED & CLEAN VERSION ===
            completed_trades = df[df['Trade Type'] == 'COMPLETED']
            if not completed_trades.empty:
                def count_outcomes(group):
                    winners = (group['Outcome'] == 'WINNER').sum()
                    losers = (group['Outcome'] == 'LOSER').sum()
                    breakeven = (group['Outcome'] == 'BREAKEVEN').sum()
                    total = len(group)
                    win_rate = (winners / total * 100) if total > 0 else 0

                    return pd.Series({
                        'Trade_Count': total,
                        'Winners': int(winners),
                        'Losers': int(losers),
                        'Breakeven': int(breakeven),
                        'Win_Rate_%': round(win_rate, 1)
                    })

                # Fixed version – clean and no warning/error
                completed_summary = (
                    completed_trades
                    .groupby(['Symbol', 'Timeframe', 'Direction', 'Order Type'], as_index=False)
                    .apply(count_outcomes, include_groups=False)
                )

                completed_summary = completed_summary.sort_values(
                    by=['Symbol', 'Timeframe', 'Direction', 'Order Type', 'Trade_Count'],
                    ascending=[True, True, True, True, False]
                ).reset_index(drop=True)
                completed_summary.to_excel(writer, sheet_name="Completed_Trades_Summary", index=False)

                # Additional detailed breakdown: Symbol + Timeframe + Direction (regardless of order type)
                                # Symbol + Timeframe + Direction sheet
                            # === Symbol + Timeframe + Direction sheet (NO Total/Avg Profit) ===
            stfd_stats = completed_trades.groupby(['Symbol', 'Timeframe', 'Direction']).agg({
                'Profit': 'count',                     # only count of trades
                'Outcome': lambda x: (x == 'WINNER').mean() * 100
            }).round(2)
            stfd_stats.columns = ['Trade_Count', 'Win_Rate_%']
            stfd_stats = stfd_stats.reset_index()

            # Alphabetical order
            stfd_stats = stfd_stats.sort_values(
                by=['Symbol', 'Timeframe', 'Direction'],
                ascending=True
            ).reset_index(drop=True)

            stfd_stats.to_excel(writer, sheet_name="Symbol_Timeframe_Direction", index=False)
            
            # === 3. STRATEGY RECOMMENDATIONS (already includes Direction) ===
                        # === 3. STRATEGY RECOMMENDATIONS (without Total/Avg Profit) ===
            strategy_data = []
            if not completed_trades.empty:
                perf = completed_trades.groupby(['Symbol', 'Timeframe', 'Direction']).agg({
                    'Profit': 'count',
                    'Outcome': lambda x: (x == 'WINNER').mean() * 100
                }).round(2)
                perf.columns = ['Trade_Count', 'Win_Rate_%']
                perf = perf.reset_index()

                for _, row in perf.iterrows():
                    recommendation = "EXCELLENT" if (row['Win_Rate_%'] >= 60 and row['Trade_Count'] >= 5) \
                        else "GOOD" if (row['Win_Rate_%'] >= 55) \
                        else "NEEDS REVIEW" if (row['Win_Rate_%'] < 45) \
                        else "MONITOR"
                    
                    strategy_data.append({
                        'Symbol': row['Symbol'],
                        'Timeframe': row['Timeframe'],
                        'Direction': row['Direction'],
                        'Trade_Count': row['Trade_Count'],
                        'Win_Rate_%': row['Win_Rate_%'],
                        'Recommendation': recommendation
                    })
            
            strategy_df = pd.DataFrame(strategy_data)
            strategy_df = strategy_df.sort_values(
                by=['Recommendation', 'Symbol', 'Timeframe', 'Direction'],
                ascending=[False, True, True, True]
            ).reset_index(drop=True)

            strategy_df.to_excel(writer, sheet_name="Strategy_Recommendations", index=False)
            
            # === Apply nice formatting ===
            workbook = writer.book
            header_font = Font(bold=True, color="FFFFFF", size=12)
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                                 top=Side(style='thin'), bottom=Side(style='thin'))

            for sheet_name in writer.sheets:
                ws = writer.sheets[sheet_name]
                # Header styling
                for cell in ws[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                
                # Auto-adjust column widths
                for col in ws.columns:
                    max_length = 0
                    column = col[0].column_letter
                    for cell in col:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    ws.column_dimensions[column].width = adjusted_width
                
                # Add borders to all cells
                for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                    for cell in row:
                        cell.border = thin_border

        print(f"Comprehensive report with DIRECTION analysis saved: {filename}")
        
    except Exception as e:
        print(f"Error saving Excel file: {e}")
        import traceback
        traceback.print_exc()

def main():
    """Main analysis function"""
    print("🚀 Starting Comprehensive Deriv Bot Analysis...")
    clear_output_folder()
    
    # Run comprehensive analysis
    df = analyze_strategy_performance()
    
    if df.empty:
        print("⚠️ No trading data found for analysis")
        mt5.shutdown()
        return
    
    # Generate reports and visualizations
    generate_comprehensive_report(df)
    create_enhanced_visualizations(df)
    save_comprehensive_excel_report(df)
    
    mt5.shutdown()
    print(f"\n✅ Comprehensive analysis completed successfully!")
    print(f"📁 Results saved to: {OUTPUT_FOLDER}")

if __name__ == "__main__":
    main()